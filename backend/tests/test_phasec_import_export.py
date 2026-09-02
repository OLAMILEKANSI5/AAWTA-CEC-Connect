"""Phase C — CSV/XLSX registration import, bulk export (segments/fields/formats/contact
gating) and bulk actions. Run seed_badge.py, seed_phasea.py, seed_phasec.py first."""
import csv
import io
import os
import uuid

import pytest
import requests
from dotenv import dotenv_values
from openpyxl import Workbook

frontend_env = dotenv_values("/app/frontend/.env")
base_url = os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")
if not base_url:
    raise RuntimeError("REACT_APP_BACKEND_URL missing")
BASE_URL = base_url.rstrip("/")
API = f"{BASE_URL}/api"

ADMIN = "tok_admtest"
MEMBER = "tok_memba"
EVENT = "evt_246d20a66b31"
OWNER_EVENT = "evt_phasec_owner"  # owned by user_memba (non-staff owner)
TAG = uuid.uuid4().hex[:6]

IMPORT_FIELDS = ["first_name", "last_name", "email", "phone", "state", "city", "zip_code", "member_status",
                 "business_name", "business_website", "job_title", "ticket_category", "payment_status",
                 "badge_category", "newsletter_opt_in", "events_opt_in", "internal_notes", "confirmation_number"]
PAYMENT_STATUSES = ["paid", "pending", "complimentary", "refunded", "not_applicable"]

HEADERS = ["First Name", "Last Name", "Email", "Phone", "State", "City", "Zip Code", "AAWTA Member",
           "Business Name", "Website", "Job Title", "Category", "Payment Status", "Badge Category",
           "Newsletter", "Upcoming Events", "Notes", "Confirmation Number"]


def _rows():
    """5 rows: 2 fully valid, 1 missing last name, 1 bad email, 1 short phone."""
    return [
        ["Alpha", f"TESTC{TAG}", f"alpha.{TAG}@example.com", "555-201-0001", "Georgia", "Atlanta", "30301",
         "Yes", "Alpha Biz", "alphabiz.com", "Owner", "VIP", "paid", "Member", "Yes", "1", "note A", ""],
        ["Bravo", f"TESTC{TAG}", f"bravo.{TAG}@example.com", "555-201-0002", "Alabama", "Mobile", "36602",
         "No", "Bravo LLC", "https://bravo.co", "Manager", "General Admission", "bogus_status", "", "No", "False", "", ""],
        ["Charlie", "", f"charlie.{TAG}@example.com", "555-201-0003", "Georgia", "Macon", "31201",
         "No", "", "", "", "", "", "", "", "", "", ""],
        ["Delta", f"TESTC{TAG}", "delta-at-example", "555-201-0004", "Georgia", "Rome", "30161",
         "No", "", "", "", "", "", "", "", "", "", ""],
        ["Echo", f"TESTC{TAG}", f"echo.{TAG}@example.com", "12345", "Georgia", "Athens", "30601",
         "Yes", "Echo Co", "", "", "", "", "", "true", "yes", "", ""],
    ]


def _csv_bytes():
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEADERS)
    for r in _rows():
        w.writerow(r)
    return buf.getvalue().encode()


def _xlsx_bytes():
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in _rows():
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def sess(token=None):
    s = requests.Session()
    if token:
        s.headers["Authorization"] = f"Bearer {token}"
    return s


@pytest.fixture(scope="session")
def admin():
    s = sess(ADMIN)
    r = s.get(f"{API}/auth/me")
    if r.status_code != 200:
        pytest.fail(f"admin token invalid: {r.status_code} {r.text[:200]}")
    return s


@pytest.fixture(scope="session")
def member():
    s = sess(MEMBER)
    r = s.get(f"{API}/auth/me")
    if r.status_code != 200:
        pytest.fail(f"member token invalid: {r.status_code} {r.text[:200]}")
    return s


@pytest.fixture(scope="session")
def anon():
    return sess()


@pytest.fixture(scope="session")
def parsed(admin):
    r = admin.post(f"{API}/events/{EVENT}/import/parse",
                   files={"file": ("phasec.csv", _csv_bytes(), "text/csv")})
    assert r.status_code == 200, r.text[:400]
    return r.json()


# ---------- import template ----------
class TestImportTemplate:
    def test_template_header_matches_import_fields(self, admin):
        r = admin.get(f"{API}/events/{EVENT}/import-template.csv")
        assert r.status_code == 200
        assert "text/csv" in r.headers.get("content-type", "")
        rows = list(csv.reader(io.StringIO(r.text)))
        assert rows[0] == IMPORT_FIELDS
        assert len(rows[1]) == len(IMPORT_FIELDS)  # sample row

    def test_template_authz(self, member, anon):
        assert member.get(f"{API}/events/{EVENT}/import-template.csv").status_code == 403
        assert anon.get(f"{API}/events/{EVENT}/import-template.csv").status_code == 401


# ---------- parse (CSV + XLSX) ----------
class TestImportParse:
    def test_parse_csv(self, parsed):
        assert parsed["count"] == 5
        assert len(parsed["rows"]) == 5
        assert parsed["columns"] == HEADERS
        m = parsed["mapping"]
        assert m["first_name"] == "First Name"
        assert m["last_name"] == "Last Name"
        assert m["email"] == "Email"
        assert m["phone"] == "Phone"
        assert m["zip_code"] == "Zip Code"
        assert m["business_website"] == "Website"
        assert m["ticket_category"] == "Category"
        assert m["payment_status"] == "Payment Status"
        assert m["newsletter_opt_in"] == "Newsletter"
        assert m["events_opt_in"] == "Upcoming Events"
        assert m["internal_notes"] == "Notes"
        assert m["confirmation_number"] == "Confirmation Number"

    def test_parse_xlsx(self, admin):
        r = admin.post(f"{API}/events/{EVENT}/import/parse",
                       files={"file": ("phasec.xlsx", _xlsx_bytes(),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r.status_code == 200, r.text[:400]
        d = r.json()
        assert d["count"] == 5, d
        assert d["columns"] == HEADERS
        assert d["mapping"]["first_name"] == "First Name"
        assert d["rows"][0]["Email"] == f"alpha.{TAG}@example.com"

    def test_parse_authz(self, member, anon):
        f = {"file": ("phasec.csv", _csv_bytes(), "text/csv")}
        assert member.post(f"{API}/events/{EVENT}/import/parse", files=f).status_code == 403
        assert anon.post(f"{API}/events/{EVENT}/import/parse",
                         files={"file": ("phasec.csv", _csv_bytes(), "text/csv")}).status_code == 401


# ---------- preview ----------
class TestImportPreview:
    def test_preview_validation_report(self, admin, parsed):
        r = admin.post(f"{API}/events/{EVENT}/import/preview",
                       json={"mapping": parsed["mapping"], "rows": parsed["rows"]})
        assert r.status_code == 200, r.text[:400]
        d = r.json()
        assert d["total"] == 5
        assert d["missing_required"] == 1
        assert d["invalid_emails"] == 1
        assert d["invalid_phones"] == 1
        assert d["skipped"] == 2
        assert d["duplicates"] == 0
        assert d["valid"] == 3
        recs = {rec["row"]: rec for rec in d["records"]}
        assert "missing_required" in recs[3]["issues"]
        assert "invalid_email" in recs[4]["issues"]
        assert "invalid_phone" in recs[5]["issues"]
        assert recs[1]["issues"] == []
        assert all(rec["is_duplicate"] is False for rec in d["records"])

    def test_preview_duplicate_detection_against_existing(self, admin):
        existing = admin.get(f"{API}/events/{EVENT}/registrations").json()["registrations"]
        target = next((x for x in existing if x.get("email")), None)
        assert target, "no existing registration with email to test duplicates"
        rows = [{"First Name": target.get("first_name") or "Dup", "Last Name": target.get("last_name") or "Dup",
                 "Email": target["email"].upper()}]
        r = admin.post(f"{API}/events/{EVENT}/import/preview",
                       json={"mapping": {"first_name": "First Name", "last_name": "Last Name", "email": "Email"},
                             "rows": rows})
        assert r.status_code == 200
        d = r.json()
        assert d["duplicates"] == 1
        rec = d["records"][0]
        assert rec["is_duplicate"] is True
        assert rec["duplicate_reason"] == "email"
        assert rec["duplicate_id"] == target["id"]

    def test_preview_duplicate_by_confirmation_number(self, admin):
        existing = admin.get(f"{API}/events/{EVENT}/registrations").json()["registrations"]
        target = next((x for x in existing if x.get("confirmation_number")), None)
        assert target
        rows = [{"First Name": "Zeta", "Last Name": "Conf", "Confirmation Number": target["confirmation_number"]}]
        r = admin.post(f"{API}/events/{EVENT}/import/preview",
                       json={"mapping": {"first_name": "First Name", "last_name": "Last Name",
                                         "confirmation_number": "Confirmation Number"}, "rows": rows})
        d = r.json()
        assert d["records"][0]["duplicate_reason"] == "confirmation_number", d["records"][0]

    def test_preview_authz(self, member, anon, parsed):
        body = {"mapping": parsed["mapping"], "rows": parsed["rows"]}
        assert member.post(f"{API}/events/{EVENT}/import/preview", json=body).status_code == 403
        assert anon.post(f"{API}/events/{EVENT}/import/preview", json=body).status_code == 401


# ---------- commit ----------
class TestImportCommit:
    def test_commit_skip_inserts_valid_rows(self, admin, parsed):
        r = admin.post(f"{API}/events/{EVENT}/import/commit",
                       json={"mapping": parsed["mapping"], "rows": parsed["rows"], "duplicate_strategy": "skip"})
        assert r.status_code == 200, r.text[:400]
        d = r.json()
        assert d["imported"] == 3, d
        assert d["failed"] == 2, d
        assert d["duplicates"] == 0
        reasons = sorted(e["reason"] for e in d["errors"])
        assert reasons == ["invalid email", "missing first/last name"], d["errors"]

        regs = admin.get(f"{API}/events/{EVENT}/registrations", params={"q": f"TESTC{TAG}"}).json()["registrations"]
        alpha = next(x for x in regs if x["first_name"] == "Alpha")
        assert alpha["source"] == "Bulk Import"
        assert alpha["confirmation_number"] and alpha["checkin_token"]
        assert alpha["is_member"] is True
        assert alpha["business_website"] == "https://alphabiz.com"
        assert alpha["newsletter_opt_in"] is True and alpha["events_opt_in"] is True
        assert alpha["ticket_category"] == "VIP"
        assert alpha["payment_status"] == "paid"
        assert alpha["name"] == f"Alpha TESTC{TAG}"
        bravo = next(x for x in regs if x["first_name"] == "Bravo")
        assert bravo["payment_status"] == "not_applicable"  # invalid value coerced
        assert bravo["is_member"] is False
        assert bravo["newsletter_opt_in"] is False and bravo["events_opt_in"] is False
        assert bravo["badge_category"] == "General Attendee"
        echo = next(x for x in regs if x["first_name"] == "Echo")
        assert echo["badge_category"] == "Member"
        assert echo["newsletter_opt_in"] is True

    def test_rerun_same_file_skip_marks_duplicates(self, admin, parsed):
        r = admin.post(f"{API}/events/{EVENT}/import/commit",
                       json={"mapping": parsed["mapping"], "rows": parsed["rows"], "duplicate_strategy": "skip"})
        d = r.json()
        assert d["imported"] == 0, d
        assert d["duplicates"] == 3, d
        assert d["skipped"] == 3, d
        assert d["failed"] == 2

    def test_update_strategy_updates_existing(self, admin, parsed):
        rows = [dict(parsed["rows"][0])]
        rows[0]["Job Title"] = "Chief TESTC"
        rows[0]["City"] = "Savannah"
        r = admin.post(f"{API}/events/{EVENT}/import/commit",
                       json={"mapping": parsed["mapping"], "rows": rows, "duplicate_strategy": "update"})
        d = r.json()
        assert d["updated"] == 1, d
        assert d["imported"] == 0
        regs = admin.get(f"{API}/events/{EVENT}/registrations", params={"q": f"alpha.{TAG}"}).json()["registrations"]
        assert len(regs) == 1, regs
        assert regs[0]["job_title"] == "Chief TESTC"
        assert regs[0]["city"] == "Savannah"

    def test_new_strategy_inserts_duplicate_anyway(self, admin, parsed):
        rows = [dict(parsed["rows"][0])]
        r = admin.post(f"{API}/events/{EVENT}/import/commit",
                       json={"mapping": parsed["mapping"], "rows": rows, "duplicate_strategy": "new"})
        d = r.json()
        assert d["imported"] == 1, d
        assert d["duplicates"] == 1, d
        regs = admin.get(f"{API}/events/{EVENT}/registrations", params={"q": f"alpha.{TAG}"}).json()["registrations"]
        assert len(regs) == 2, f"expected duplicate insert, got {len(regs)}"

    def test_commit_audit_logged(self, admin):
        logs = admin.get(f"{API}/audit-logs", params={"q": "registrations_imported"}).json()
        assert any(x["action"] == "registrations_imported" and x["related_id"] == EVENT for x in logs), logs[:3]

    def test_commit_authz(self, member, anon, parsed):
        body = {"mapping": parsed["mapping"], "rows": [], "duplicate_strategy": "skip"}
        assert member.post(f"{API}/events/{EVENT}/import/commit", json=body).status_code == 403
        assert anon.post(f"{API}/events/{EVENT}/import/commit", json=body).status_code == 401


# ---------- export ----------
def _csv_rows(resp):
    return list(csv.reader(io.StringIO(resp.text)))


class TestExport:
    def test_csv_format_and_headers(self, admin):
        r = admin.get(f"{API}/events/{EVENT}/export", params={"format": "csv"})
        assert r.status_code == 200
        assert "text/csv" in r.headers["content-type"]
        rows = _csv_rows(r)
        assert rows[0][0] == "Name"
        assert "Email" in rows[0] and "Phone" in rows[0]
        assert len(rows) > 1

    def test_xlsx_and_pdf_content_types(self, admin):
        rx = admin.get(f"{API}/events/{EVENT}/export", params={"format": "xlsx"})
        assert rx.status_code == 200
        assert rx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert rx.content[:2] == b"PK"
        rp = admin.get(f"{API}/events/{EVENT}/export", params={"format": "pdf"})
        assert rp.status_code == 200
        assert rp.headers["content-type"] == "application/pdf"
        assert rp.content[:4] == b"%PDF"

    def test_fields_param_limits_and_orders(self, admin):
        r = admin.get(f"{API}/events/{EVENT}/export",
                      params={"format": "csv", "fields": "email,name,state"})
        rows = _csv_rows(r)
        assert rows[0] == ["Email", "Name", "State"]

    def test_ids_param_exports_only_selected(self, admin):
        regs = admin.get(f"{API}/events/{EVENT}/registrations").json()["registrations"]
        ids = [x["id"] for x in regs[:2]]
        r = admin.get(f"{API}/events/{EVENT}/export",
                      params={"format": "csv", "ids": ",".join(ids), "fields": "name"})
        rows = _csv_rows(r)
        assert len(rows) == 3, rows  # header + 2
        names = {x["name"] for x in regs[:2]}
        assert {rows[1][0], rows[2][0]} == names

    @pytest.mark.parametrize("segment", ["all", "checked_in", "not_checked_in", "walkins", "manual",
                                         "online", "imported", "members", "nonmembers", "newsletter", "events"])
    def test_segments_filter(self, admin, segment):
        r = admin.get(f"{API}/events/{EVENT}/export", params={"format": "csv", "segment": segment})
        assert r.status_code == 200, r.text[:200]
        rows = _csv_rows(r)
        exported = len(rows) - 1
        regs = admin.get(f"{API}/events/{EVENT}/registrations").json()["registrations"]
        pred = {
            "all": lambda x: True,
            "checked_in": lambda x: bool(x.get("checked_in")),
            "not_checked_in": lambda x: not x.get("checked_in"),
            "walkins": lambda x: x.get("source") == "Walk-In",
            "manual": lambda x: x.get("source") == "Manual Entry",
            "online": lambda x: x.get("source") == "online",
            "imported": lambda x: x.get("source") == "Bulk Import",
            "members": lambda x: bool(x.get("is_member")),
            "nonmembers": lambda x: not x.get("is_member"),
            "newsletter": lambda x: bool(x.get("newsletter_opt_in")),
            "events": lambda x: bool(x.get("events_opt_in")),
        }[segment]
        assert exported == len([x for x in regs if pred(x)]), f"segment {segment} mismatch"

    def test_state_and_category_filters(self, admin):
        r = admin.get(f"{API}/events/{EVENT}/export",
                      params={"format": "csv", "state": "Georgia", "fields": "name,state"})
        rows = _csv_rows(r)
        assert all(row[1] == "Georgia" for row in rows[1:])
        r2 = admin.get(f"{API}/events/{EVENT}/export",
                       params={"format": "csv", "category": "VIP", "fields": "name,ticket_category"})
        rows2 = _csv_rows(r2)
        assert all(row[1] == "VIP" for row in rows2[1:])

    def test_staff_export_includes_contact_fields(self, admin):
        r = admin.get(f"{API}/events/{OWNER_EVENT}/export", params={"format": "csv"})
        assert r.status_code == 200, r.text[:200]
        rows = _csv_rows(r)
        assert "Email" in rows[0] and "Phone" in rows[0]
        assert any("ownergate@example.com" in c for c in rows[1])

    def test_nonstaff_owner_export_strips_contact_fields(self, member):
        r = member.get(f"{API}/events/{OWNER_EVENT}/export", params={"format": "csv"})
        assert r.status_code == 200, r.text[:300]
        rows = _csv_rows(r)
        assert "Email" not in rows[0], rows[0]
        assert "Phone" not in rows[0], rows[0]
        assert "ownergate@example.com" not in r.text
        assert "555-000-1111" not in r.text
        # explicit fields request must also be stripped
        r2 = member.get(f"{API}/events/{OWNER_EVENT}/export",
                        params={"format": "csv", "fields": "name,email,phone"})
        rows2 = _csv_rows(r2)
        assert rows2[0] == ["Name"], rows2[0]

    def test_nonstaff_owner_xlsx_pdf_strip_contacts(self, member):
        rp = member.get(f"{API}/events/{OWNER_EVENT}/export", params={"format": "pdf"})
        assert rp.status_code == 200
        assert rp.headers["content-type"] == "application/pdf"
        rx = member.get(f"{API}/events/{OWNER_EVENT}/export", params={"format": "xlsx"})
        assert rx.status_code == 200
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(rx.content))
        header = [c.value for c in next(wb.active.iter_rows())]
        assert "Email" not in header and "Phone" not in header, header

    def test_export_authz(self, member, anon):
        assert member.get(f"{API}/events/{EVENT}/export").status_code == 403
        assert anon.get(f"{API}/events/{EVENT}/export").status_code == 401

    def test_export_audit_logged(self, admin):
        logs = admin.get(f"{API}/audit-logs", params={"q": "registrations_exported"}).json()
        assert any(x["action"] == "registrations_exported" for x in logs)


# ---------- bulk actions ----------
@pytest.fixture(scope="class")
def bulk_targets(admin):
    """Two imported TESTC rows used as bulk-action targets."""
    regs = admin.get(f"{API}/events/{EVENT}/registrations", params={"q": f"TESTC{TAG}"}).json()["registrations"]
    ids = [x["id"] for x in regs][:2]
    assert len(ids) == 2, f"expected 2 imported targets, got {len(ids)}"
    return ids


def _get(admin, reg_id, include_archived=False):
    regs = admin.get(f"{API}/events/{EVENT}/registrations",
                     params={"include_archived": str(include_archived).lower()}).json()["registrations"]
    return next((x for x in regs if x["id"] == reg_id), None)


class TestBulkActions:
    def test_change_category(self, admin, bulk_targets):
        r = admin.post(f"{API}/events/{EVENT}/bulk",
                       json={"action": "change_category", "ids": bulk_targets, "value": "VIP"})
        assert r.status_code == 200
        d = r.json()
        assert d == {"ok": True, "action": "change_category", "affected": 2}
        for i in bulk_targets:
            assert _get(admin, i)["ticket_category"] == "VIP"

    def test_assign_badge_category(self, admin, bulk_targets):
        r = admin.post(f"{API}/events/{EVENT}/bulk",
                       json={"action": "assign_badge_category", "ids": bulk_targets, "value": "Speaker"})
        assert r.json()["affected"] == 2
        for i in bulk_targets:
            assert _get(admin, i)["badge_category"] == "Speaker"

    def test_mark_complimentary(self, admin, bulk_targets):
        r = admin.post(f"{API}/events/{EVENT}/bulk",
                       json={"action": "mark_complimentary", "ids": bulk_targets})
        assert r.json()["affected"] == 2
        for i in bulk_targets:
            assert _get(admin, i)["payment_status"] == "complimentary"

    def test_update_payment_valid_and_invalid(self, admin, bulk_targets):
        r = admin.post(f"{API}/events/{EVENT}/bulk",
                       json={"action": "update_payment", "ids": bulk_targets, "value": "refunded"})
        assert r.json()["affected"] == 2
        for i in bulk_targets:
            assert _get(admin, i)["payment_status"] == "refunded"
        bad = admin.post(f"{API}/events/{EVENT}/bulk",
                         json={"action": "update_payment", "ids": bulk_targets, "value": "not_a_status"})
        assert bad.status_code == 200
        assert bad.json()["affected"] == 0, "invalid payment status must not update rows"
        for i in bulk_targets:
            assert _get(admin, i)["payment_status"] == "refunded"

    def test_resend_confirmation_bumps_counter(self, admin, bulk_targets):
        before = {i: (_get(admin, i).get("resend_count") or 0) for i in bulk_targets}
        r = admin.post(f"{API}/events/{EVENT}/bulk",
                       json={"action": "resend_confirmation", "ids": bulk_targets})
        assert r.json()["affected"] == 2
        for i in bulk_targets:
            after = _get(admin, i)
            assert (after.get("resend_count") or 0) == before[i] + 1
            assert after.get("last_resend_at")

    def test_checkin(self, admin, bulk_targets):
        r = admin.post(f"{API}/events/{EVENT}/bulk", json={"action": "checkin", "ids": bulk_targets})
        assert r.json()["affected"] == 2
        for i in bulk_targets:
            reg = _get(admin, i)
            assert reg["checked_in"] is True
            assert reg["status"] == "checked_in"
            assert reg["checkin_method"] == "bulk"

    def test_cancel(self, admin, bulk_targets):
        r = admin.post(f"{API}/events/{EVENT}/bulk", json={"action": "cancel", "ids": bulk_targets})
        assert r.json()["affected"] == 2
        for i in bulk_targets:
            reg = _get(admin, i)
            assert reg["status"] == "cancelled"
            assert reg["checked_in"] is False

    def test_archive_hides_from_default_list(self, admin, bulk_targets):
        r = admin.post(f"{API}/events/{EVENT}/bulk", json={"action": "archive", "ids": bulk_targets})
        assert r.json()["affected"] == 2
        for i in bulk_targets:
            assert _get(admin, i) is None
            assert _get(admin, i, include_archived=True)["archived"] is True

    def test_unknown_action_affects_nothing(self, admin, bulk_targets):
        r = admin.post(f"{API}/events/{EVENT}/bulk", json={"action": "nope", "ids": bulk_targets})
        assert r.status_code == 200
        assert r.json()["affected"] == 0

    def test_bulk_audit_entries(self, admin):
        logs = admin.get(f"{API}/audit-logs", params={"q": "bulk_"}).json()
        actions = {x["action"] for x in logs}
        for a in ("bulk_change_category", "bulk_assign_badge_category", "bulk_mark_complimentary",
                  "bulk_update_payment", "bulk_resend_confirmation", "bulk_checkin", "bulk_cancel", "bulk_archive"):
            assert a in actions, f"missing audit entry {a} (found {sorted(actions)})"

    def test_bulk_authz(self, member, anon):
        body = {"action": "checkin", "ids": []}
        assert member.post(f"{API}/events/{EVENT}/bulk", json=body).status_code == 403
        assert anon.post(f"{API}/events/{EVENT}/bulk", json=body).status_code == 401

    def test_nonstaff_owner_can_bulk_on_own_event(self, member):
        r = member.post(f"{API}/events/{OWNER_EVENT}/bulk",
                        json={"action": "change_category", "ids": ["reg_phasec_owner1"], "value": "General Admission"})
        assert r.status_code == 200, r.text[:200]
        assert r.json()["affected"] == 1

    def test_missing_event_404(self, admin):
        assert admin.post(f"{API}/events/evt_does_not_exist/bulk",
                          json={"action": "checkin", "ids": []}).status_code == 404
        assert admin.get(f"{API}/events/evt_does_not_exist/export").status_code == 404
